#!/usr/bin/env python3
"""
Converte uma aba de mesociclo da planilha (ex.: "M (6)") para o JSON que o workflow
"Importar mesociclo" grava no Firestore: prescrição (mesmo formato de mesoParaDoc) + cargas já
registradas. Precisa de openpyxl (pip install openpyxl). O .xlsx nunca vai para o repositório.

Uso: python3 tools/planilha_para_meso.py planilha.xlsx "M (6)" [--frequencia N] > meso.json
"""
import argparse
import json
import sys
import unicodedata
import uuid
from datetime import datetime

import openpyxl
from openpyxl.utils import column_index_from_string as col

METODO_COLS = ['id', 'ativo', 'semana', 'treino', 'exercicio', 'nomeExercicio', 'metodo', 'escopo', 'serie', 'grupo',
               'reducaoCarga', 'pausaInternaS', 'blocos', 'repsPorBloco', 'adicional', 'etapas',
               'cargaEtapa1', 'cargaEtapa2', 'cargaEtapa3', 'repsPadrao', 'repsEtapa1', 'repsEtapa2', 'repsEtapa3']


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def texto(v):
    return v.strip() if isinstance(v, str) else ''


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('aba')
    ap.add_argument('--frequencia', type=int, help='sobrescreve "Sessões [nx]" (AZ5 ou BO7)')
    a = ap.parse_args()
    ws = openpyxl.load_workbook(a.xlsx, data_only=True)[a.aba]
    c = lambda r, cl: ws.cell(r, cl if isinstance(cl, int) else col(cl)).value

    # Config: rótulo "Semana [n]" ou "Número de semanas" (sem acento/maiúsculas/espaços extras) em AY3
    # (abas antigas) ou BN5 (abas novas, com o monitoramento em AY);
    # valores na coluna ao lado; a tabela de métodos começa 7 linhas abaixo do rótulo.
    rotulo = lambda v: ' '.join(''.join(ch for ch in unicodedata.normalize('NFD', texto(v)) if not unicodedata.combining(ch)).lower().split())
    cr, cc = next(((r, k) for r in range(1, 11) for k in range(1, ws.max_column + 1)
                   if rotulo(c(r, k)) in ('semana [n]', 'numero de semanas')), (3, col('AY')))
    semanas = int(c(cr, cc + 1))
    inicio = c(cr + 1, cc + 1)
    if not isinstance(inicio, datetime):
        sys.exit(f'Data inicial ({ws.cell(cr + 1, cc + 1).coordinate}) não é uma data: {inicio!r}')
    frequencia = a.frequencia or int(c(cr + 2, cc + 1))
    pausa = int(num(c(33, 'Z')) or 120)

    # Divisão: T4:Y13 (coluna = treino, linha = posição).
    divisao = []
    for t in range(frequencia):
        slots = []
        for p in range(10):
            nome = texto(c(4 + p, col('T') + t))
            if not nome:
                break
            slots.append({'slotId': str(uuid.uuid4()), 'nome': nome})
        divisao.append(slots)

    def linha(r):
        vals = [num(c(r, col('T') + s)) for s in range(6)]
        return [int(v) if v is not None else None for v in vals]

    progressao = {'principais': [linha(18 + w) for w in range(semanas)],
                  'auxiliares': [linha(33 + w) for w in range(semanas)]}

    # Métodos: tabela AY10:BU109 (abas novas: BN12:CJ111). Linhas repetidas (mesmo conteúdo) entram uma vez.
    metodos, vistos = [], set()
    for r in range(cr + 7, cr + 107):
        m = dict(zip(METODO_COLS, (c(r, cc + i) for i in range(len(METODO_COLS)))))
        if not texto(m['metodo']) or not num(m['semana']) or not num(m['treino']) or not num(m['exercicio']):
            continue
        if m['treino'] > frequencia:
            continue
        ap_ = {'ativo': texto(m['ativo']).upper() == 'SIM', 'semana': int(m['semana']), 'treino': int(m['treino']),
               'exercicio': int(m['exercicio']), 'metodo': texto(m['metodo']), 'escopo': texto(m['escopo'])}
        for k in ['serie', 'grupo', 'pausaInternaS', 'blocos', 'repsPorBloco', 'adicional', 'etapas', 'repsPadrao']:
            if num(m[k]):
                ap_[k] = m[k]
        if num(m['reducaoCarga']):  # planilha: fração (0,2); app: % (20)
            ap_['reducaoCarga'] = round(m['reducaoCarga'] * 100)
        for k in ['cargaEtapa', 'repsEtapa']:
            vals = [num(m[f'{k}{i}']) for i in (1, 2, 3)]
            if any(vals):
                ap_[k] = vals
        chave = json.dumps(ap_, sort_keys=True)
        if chave not in vistos:
            vistos.add(chave)
            metodos.append(ap_)

    # Semanas: blocos de 1 + 13·n linhas, n = treinos do modelo da aba (79 no "6x", 53 no "4x"); acha pela
    # 2ª linha "Semana" em B. Treinos: blocos de 13 (cabeçalho "Realizado" em C).
    # Nome diferente da divisão numa semana = troca. Cargas em J:O (séries 1–6).
    bloco = next((r for r in range(3, ws.max_row + 1) if texto(c(r, 'B')) == 'Semana'), 81) - 2
    trocas, registros = {}, []
    for w in range(1, semanas + 1):
        for t in range(1, frequencia + 1):
            topo = 3 + bloco * (w - 1) + 13 * (t - 1)
            if texto(c(topo, 'C')) not in ('Realizado', 'Não realizado'):
                sys.exit(f'Layout inesperado na semana {w}, treino {t} (linha {topo})')
            cargas = {}
            for p, slot in enumerate(divisao[t - 1], start=1):
                nome = texto(c(topo + 1 + p, 'C'))
                if nome and nome != slot['nome']:
                    trocas[f'{w}-{t}-{p}'] = slot = {'slotId': str(uuid.uuid4()), 'nome': nome}
                kg = {f's{s + 1}': v for s in range(6) if (v := num(c(topo + 1 + p, col('J') + s)))}  # 0 = vazio
                if kg:
                    cargas[slot['slotId']] = kg
            realizado = c(topo, 'B') is True
            if cargas or realizado:
                registros.append({'semana': w, 'treino': t, 'realizado': realizado, 'cargas': cargas})

    json.dump({
        'meso': {
            'frequencia': frequencia, 'semanas': semanas, 'inicio': inicio.strftime('%Y-%m-%d'), 'pausaS': pausa,
            'divisao': {f't{i + 1}': s for i, s in enumerate(divisao)},
            'trocas': trocas,
            'progressao': {k: {f'w{i + 1}': l for i, l in enumerate(v)} for k, v in progressao.items()},
            'metodos': metodos,
        },
        'registros': registros,
    }, sys.stdout, ensure_ascii=False, separators=(',', ':'))


if __name__ == '__main__':
    main()
